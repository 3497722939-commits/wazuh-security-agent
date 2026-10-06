"""把毕业论文.md 转成格式规范的 docx。"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = r"D:\agent\毕业论文.md"
DST = r"D:\agent\毕业论文.docx"

doc = Document()

# 页面设置
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

# 默认字体
style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(12)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

def set_cn_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)

def add_heading(text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 1:
        set_cn_font(run, "黑体", 16, True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(12)
    elif level == 2:
        set_cn_font(run, "黑体", 14, True)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    elif level == 3:
        set_cn_font(run, "黑体", 12, True)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
    return p

def add_body(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cn_font(run, "宋体", 12)
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    return p

def add_table(rows):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = table.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            run = p.add_run(cell.strip())
            set_cn_font(run, "宋体", 10, bold=(i == 0))
    return table

# 读取 markdown
with open(SRC, encoding="utf-8") as f:
    lines = f.readlines()

i = 0
while i < len(lines):
    line = lines[i].rstrip()

    # 跳过空行
    if not line.strip():
        i += 1
        continue

    # 分隔线
    if line.strip() == "---":
        i += 1
        continue

    # 标题
    m = re.match(r"^(#{1,3})\s+(.*)", line)
    if m:
        level = len(m.group(1))
        add_heading(m.group(2).strip(), level)
        i += 1
        continue

    # 表格
    if line.strip().startswith("|"):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            table_lines.append(lines[i].strip())
            i += 1
        # 解析表格
        rows = []
        for tl in table_lines:
            if re.match(r"^\|[\s\-:|]+\|$", tl):
                continue  # 分隔行
            cells = [c.strip() for c in tl.split("|")[1:-1]]
            rows.append(cells)
        add_table(rows)
        continue

    # 列表
    if re.match(r"^\d+\.\s", line.strip()) or line.strip().startswith("- "):
        text = re.sub(r"^[\d\-\*]+\s+", "", line.strip())
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_cn_font(run, "宋体", 12)
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        i += 1
        continue

    # 代码块
    if line.strip().startswith("```"):
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            code_lines.append(lines[i].rstrip())
            i += 1
        i += 1
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        p.paragraph_format.left_indent = Pt(24)
        continue

    # 普通段落
    add_body(line.strip())
    i += 1

doc.save(DST)
print(f"Saved: {DST}")
